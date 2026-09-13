from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path


WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
PYTHON = sys.executable
TEMP_ROOT = WORKSPACE_ROOT / "var" / "temp"
SPREADSHEET_TOOL = WORKSPACE_ROOT / "tools" / "extract-spreadsheet.py"
PDF_TOOL = WORKSPACE_ROOT / "tools" / "extract-pdf-pages.py"


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [PYTHON, *args],
        cwd=WORKSPACE_ROOT,
        text=True,
        encoding="utf-8",
        errors="replace",
        capture_output=True,
        check=False,
    )


class SpreadsheetCleanTests(unittest.TestCase):
    def test_clean_exports_nonempty_sheets_without_hash_binding(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(dir=TEMP_ROOT) as temporary:
            project = Path(temporary) / "project"
            raw = project / "02-data" / "raw"
            raw.mkdir(parents=True)
            source = raw / "Source Data.xlsx"

            workbook = Workbook()
            data = workbook.active
            data.title = "Data"
            data.append([None, None, None, None])
            data.append([" ID ", "Value (%)", None, "Value (%)"])
            data.append(["0012", " 3.5 ", None, date(2026, 9, 13)])
            data.append([None, None, None, None])
            workbook.create_sheet("空表")
            workbook.save(source)

            result = run(str(SPREADSHEET_TOOL), "clean", str(source))
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

            processed = project / "02-data" / "processed"
            output = processed / "source-data-sheet-01-data.csv"
            report = processed / "source-data-cleaning-report.json"
            self.assertTrue(output.is_file())
            self.assertTrue(report.is_file())
            self.assertFalse((processed / "source-data-sheet-02.csv").exists())

            with output.open(encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[0], ["id", "value", "value_2"])
            self.assertEqual(rows[1], ["0012", "3.5", "2026-09-13"])
            self.assertEqual(len(rows), 2)

            payload = json.loads(report.read_text(encoding="utf-8"))
            self.assertNotIn("sha256", report.read_text(encoding="utf-8").casefold())
            self.assertTrue(payload["input"]["exists"])
            self.assertEqual(payload["sheets"][0]["removed_fully_empty_columns"], 1)
            self.assertEqual(payload["sheets"][1]["status"], "skipped-empty")

    def test_inspect_report_uses_filename_and_existence(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory(dir=TEMP_ROOT) as temporary:
            root = Path(temporary)
            source = root / "source.xlsx"
            report = root / "inspect.json"
            workbook = Workbook()
            workbook.active.append(["x"])
            workbook.active.append([1])
            workbook.save(source)

            result = run(
                str(SPREADSHEET_TOOL),
                "inspect",
                str(source),
                "--report",
                str(report),
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            payload = json.loads(report.read_text(encoding="utf-8"))
            self.assertEqual(payload["input"]["exists"], True)
            self.assertNotIn("sha256", payload["input"])


class PdfExtractTests(unittest.TestCase):
    def test_page_and_crop_exports_pdf_png_to_temp_area(self) -> None:
        import pymupdf

        with tempfile.TemporaryDirectory(dir=TEMP_ROOT) as temporary:
            root = Path(temporary)
            source = root / "Problem.pdf"
            output_dir = root / "captures"
            document = pymupdf.open()
            for index in range(2):
                page = document.new_page(width=200, height=100)
                page.insert_text((20, 30), f"page {index + 1}")
            document.save(source)
            document.close()

            result = run(
                str(PDF_TOOL),
                str(source),
                "--pages",
                "1-2",
                "--crop",
                "0.25",
                "0.25",
                "0.75",
                "0.75",
                "--format",
                "both",
                "--dpi",
                "72",
                "--output-dir",
                str(output_dir),
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

            pdf = output_dir / "problem-pages-001-002-crop.pdf"
            png = output_dir / "problem-page-001-crop.png"
            report = output_dir / "problem-pages-001-002-crop-capture-report.json"
            self.assertTrue(pdf.is_file())
            self.assertTrue(png.is_file())
            self.assertTrue(report.is_file())

            cropped_pdf = pymupdf.open(pdf)
            try:
                self.assertEqual(cropped_pdf.page_count, 2)
                self.assertAlmostEqual(cropped_pdf[0].rect.width, 100, delta=0.5)
                self.assertAlmostEqual(cropped_pdf[0].rect.height, 50, delta=0.5)
            finally:
                cropped_pdf.close()
            pixmap = pymupdf.Pixmap(png)
            self.assertEqual((pixmap.width, pixmap.height), (100, 50))
            self.assertNotIn("sha256", report.read_text(encoding="utf-8").casefold())


if __name__ == "__main__":
    unittest.main()
