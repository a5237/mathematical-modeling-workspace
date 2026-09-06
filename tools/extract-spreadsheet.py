#!/usr/bin/env python3
"""Inspect Excel workbooks and stream worksheet data to auditable CSV/TSV files.

The source workbook is always opened read-only. Modern Excel files are streamed
with openpyxl so that large worksheets do not need to be materialized as a
DataFrame. Legacy .xls files are supported through xlrd.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import sqlite3
import sys
import tempfile
import time
from collections import Counter
from datetime import date, datetime, time as datetime_time, timezone
from pathlib import Path
from typing import Any, Iterable, Iterator, Sequence


TOOL_VERSION = "1.1.0"
WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
MODERN_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
SUPPORTED_EXCEL_SUFFIXES = MODERN_EXCEL_SUFFIXES | {".xls"}
DEFAULT_DUPLICATE_MEMORY_ROWS = 100_000


class SpreadsheetToolError(RuntimeError):
    """A user-actionable spreadsheet processing error."""


def configure_utf8_stdio() -> None:
    """Keep JSON output UTF-8 on Windows terminals and redirected streams."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is not None:
            try:
                reconfigure(encoding="utf-8", errors="replace")
            except (OSError, ValueError):
                pass


def positive_int(value: str) -> int:
    parsed = int(value)
    if parsed < 1:
        raise argparse.ArgumentTypeError("must be at least 1")
    return parsed


def nonnegative_int(value: str) -> int:
    parsed = int(value)
    if parsed < 0:
        raise argparse.ArgumentTypeError("must be at least 0")
    return parsed


def display_path(path: Path) -> str:
    """Avoid leaking personal absolute paths into reports and logs."""
    resolved = path.resolve()
    try:
        return resolved.relative_to(WORKSPACE_ROOT).as_posix()
    except ValueError:
        return path.name


def validate_input(path: Path) -> Path:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise SpreadsheetToolError(f"input file does not exist: {path}")
    if resolved.suffix.lower() not in SUPPORTED_EXCEL_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_EXCEL_SUFFIXES))
        raise SpreadsheetToolError(
            f"unsupported input format {resolved.suffix!r}; expected one of: {supported}"
        )
    return resolved


def validate_output(input_path: Path, output_path: Path, overwrite: bool) -> Path:
    resolved = output_path.expanduser().resolve()
    if resolved == input_path:
        raise SpreadsheetToolError("output path must not overwrite the source workbook")
    if resolved.exists() and not overwrite:
        raise SpreadsheetToolError(
            f"output already exists: {display_path(resolved)}; pass --overwrite explicitly"
        )
    return resolved


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return str(value)


def csv_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, datetime_time)):
        return value.isoformat()
    return value


def normalized_header(value: Any) -> str:
    return "" if value is None else str(value).strip()


def trim_trailing_empty(values: Sequence[Any]) -> list[Any]:
    result = list(values)
    while result and result[-1] is None:
        result.pop()
    return result


def header_issues(raw_headers: Sequence[Any]) -> dict[str, Any]:
    normalized = [normalized_header(value) for value in raw_headers]
    counts = Counter(value for value in normalized if value)
    return {
        "blank_column_numbers": [
            index for index, value in enumerate(normalized, start=1) if not value
        ],
        "duplicate_names": sorted(name for name, count in counts.items() if count > 1),
        "trimmed_whitespace_column_numbers": [
            index
            for index, value in enumerate(raw_headers, start=1)
            if isinstance(value, str) and value != value.strip()
        ],
    }


def repaired_headers(raw_headers: Sequence[Any]) -> tuple[list[str], list[str]]:
    headers: list[str] = []
    warnings: list[str] = []
    used: Counter[str] = Counter()
    for index, raw_value in enumerate(raw_headers, start=1):
        base = normalized_header(raw_value)
        if not base:
            base = f"column-{index}"
            warnings.append(f"blank header at column {index} renamed to {base!r}")
        elif isinstance(raw_value, str) and raw_value != raw_value.strip():
            warnings.append(f"header whitespace trimmed at column {index}")

        used[base] += 1
        name = base if used[base] == 1 else f"{base}__{used[base]}"
        if name != base:
            warnings.append(f"duplicate header {base!r} renamed to {name!r}")
        headers.append(name)
    return headers, warnings


def select_columns(
    raw_headers: Sequence[Any],
    requested: Sequence[str] | None,
    repair: bool,
) -> tuple[list[int], list[str], list[str]]:
    normalized = [normalized_header(value) for value in raw_headers]
    warnings = [
        f"header whitespace normalized at column {index}"
        for index, value in enumerate(raw_headers, start=1)
        if isinstance(value, str) and value != value.strip()
    ]

    if repair:
        effective, warnings = repaired_headers(raw_headers)
    else:
        effective = normalized

    if requested:
        if len(set(requested)) != len(requested):
            raise SpreadsheetToolError("--columns contains duplicate names")
        indices: list[int] = []
        for name in requested:
            matches = [index for index, header in enumerate(effective) if header == name]
            if not matches:
                available = ", ".join(repr(value) for value in effective if value)
                raise SpreadsheetToolError(
                    f"column {name!r} was not found; available columns: {available}"
                )
            if len(matches) > 1:
                raise SpreadsheetToolError(
                    f"column {name!r} is ambiguous; pass --repair-headers and use the repaired name"
                )
            indices.append(matches[0])
        return indices, list(requested), warnings

    issues = header_issues(raw_headers)
    if (issues["blank_column_numbers"] or issues["duplicate_names"]) and not repair:
        details: list[str] = []
        if issues["blank_column_numbers"]:
            details.append(f"blank headers at columns {issues['blank_column_numbers']}")
        if issues["duplicate_names"]:
            details.append(f"duplicate headers {issues['duplicate_names']}")
        raise SpreadsheetToolError(
            "; ".join(details) + "; select named columns or pass --repair-headers"
        )
    if not effective:
        raise SpreadsheetToolError("the selected header row is empty")
    return list(range(len(effective))), effective, warnings


def choose_sheet(
    sheet_names: Sequence[str], sheet_name: str | None, sheet_index: int | None
) -> tuple[str, int]:
    if not sheet_names:
        raise SpreadsheetToolError("the workbook contains no worksheets")
    if sheet_name is not None:
        if sheet_name not in sheet_names:
            available = ", ".join(repr(name) for name in sheet_names)
            raise SpreadsheetToolError(
                f"worksheet {sheet_name!r} was not found; available sheets: {available}"
            )
        index = sheet_names.index(sheet_name)
        return sheet_names[index], index
    if sheet_index is not None:
        index = sheet_index - 1
        if not 0 <= index < len(sheet_names):
            raise SpreadsheetToolError(
                f"--sheet-index must be between 1 and {len(sheet_names)}"
            )
        return sheet_names[index], index
    return sheet_names[0], 0


def fingerprint_row(values: Sequence[Any]) -> bytes:
    payload = json.dumps(
        [json_value(value) for value in values],
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.blake2b(payload, digest_size=16).digest()


class DuplicateTracker:
    """Count exact duplicate fingerprints with bounded automatic memory use."""

    def __init__(self, strategy: str, memory_rows: int) -> None:
        self.requested_strategy = strategy
        self.strategy_used = "off" if strategy == "off" else "memory"
        self.memory_rows = memory_rows
        self._seen: set[bytes] = set()
        self._connection: sqlite3.Connection | None = None
        self._database_path: Path | None = None
        if strategy == "disk":
            self._switch_to_disk()

    def _switch_to_disk(self) -> None:
        cache_root = WORKSPACE_ROOT / "var" / "tmp"
        cache_root.mkdir(parents=True, exist_ok=True)
        descriptor, raw_path = tempfile.mkstemp(
            prefix="spreadsheet-duplicates-", suffix=".sqlite3", dir=cache_root
        )
        os.close(descriptor)
        self._database_path = Path(raw_path)
        try:
            self._connection = sqlite3.connect(self._database_path)
            self._connection.execute("PRAGMA journal_mode=OFF")
            self._connection.execute("PRAGMA synchronous=OFF")
            self._connection.execute("PRAGMA temp_store=FILE")
            self._connection.execute(
                "CREATE TABLE fingerprints (value BLOB PRIMARY KEY) WITHOUT ROWID"
            )
            if self._seen:
                self._connection.executemany(
                    "INSERT INTO fingerprints(value) VALUES (?)",
                    ((value,) for value in self._seen),
                )
                self._seen.clear()
            self.strategy_used = "disk"
        except Exception:
            self.close()
            raise

    def observe(self, value: bytes) -> bool:
        """Return True when the fingerprint has already been observed."""
        if self.requested_strategy == "off":
            return False
        if (
            self.requested_strategy == "auto"
            and self._connection is None
            and len(self._seen) >= self.memory_rows
        ):
            self._switch_to_disk()
        if self._connection is not None:
            changes_before = self._connection.total_changes
            self._connection.execute(
                "INSERT OR IGNORE INTO fingerprints(value) VALUES (?)", (value,)
            )
            return self._connection.total_changes == changes_before
        if value in self._seen:
            return True
        self._seen.add(value)
        return False

    def close(self) -> None:
        if self._connection is not None:
            self._connection.close()
            self._connection = None
        if self._database_path is not None:
            self._database_path.unlink(missing_ok=True)
            self._database_path = None


def type_label(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, datetime_time):
        return "time"
    if isinstance(value, str):
        return "formula" if value.startswith("=") else "text"
    return type(value).__name__


def temporary_sibling(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, raw_path = tempfile.mkstemp(
        prefix=f".{path.name}.", suffix=".tmp", dir=path.parent
    )
    os.close(descriptor)
    return Path(raw_path)


def write_json_atomic(path: Path, payload: dict[str, Any]) -> None:
    temporary = temporary_sibling(path)
    try:
        temporary.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_rows(
    rows: Iterable[Sequence[Any]],
    headers: Sequence[str],
    selected_indices: Sequence[int],
    output_path: Path,
    *,
    encoding: str,
    delimiter: str,
    keep_empty_rows: bool,
    duplicate_strategy: str,
    duplicate_memory_rows: int,
    max_rows: int | None,
) -> dict[str, Any]:
    temporary = temporary_sibling(output_path)
    missing = [0 for _ in headers]
    types = [Counter() for _ in headers]
    duplicates: DuplicateTracker | None = None
    duplicate_strategy_used = "off" if duplicate_strategy == "off" else "memory"
    duplicate_rows = 0
    scanned_rows = 0
    output_rows = 0
    skipped_empty_rows = 0

    try:
        duplicates = DuplicateTracker(duplicate_strategy, duplicate_memory_rows)
        with temporary.open("w", encoding=encoding, newline="") as handle:
            writer = csv.writer(handle, delimiter=delimiter, lineterminator="\n")
            writer.writerow(headers)
            for source_row in rows:
                if max_rows is not None and output_rows >= max_rows:
                    break
                scanned_rows += 1
                selected = [
                    source_row[index] if index < len(source_row) else None
                    for index in selected_indices
                ]
                if not keep_empty_rows and all(value is None for value in selected):
                    skipped_empty_rows += 1
                    continue

                for index, value in enumerate(selected):
                    if value is None:
                        missing[index] += 1
                    else:
                        types[index][type_label(value)] += 1

                if duplicate_strategy != "off" and duplicates.observe(
                    fingerprint_row(selected)
                ):
                    duplicate_rows += 1

                writer.writerow([csv_value(value) for value in selected])
                output_rows += 1
        duplicate_strategy_used = duplicates.strategy_used
        os.replace(temporary, output_path)
    finally:
        temporary.unlink(missing_ok=True)
        if duplicates is not None:
            duplicates.close()

    return {
        "scanned_source_rows": scanned_rows,
        "output_rows": output_rows,
        "skipped_empty_rows": skipped_empty_rows,
        "duplicate_rows": duplicate_rows if duplicate_strategy != "off" else None,
        "duplicate_check_strategy_requested": duplicate_strategy,
        "duplicate_check_strategy_used": duplicate_strategy_used,
        "missing_cells_by_column": dict(zip(headers, missing, strict=True)),
        "observed_types_by_column": {
            header: dict(sorted(counter.items()))
            for header, counter in zip(headers, types, strict=True)
        },
    }


def read_modern_header(worksheet: Any, header_row: int) -> list[Any]:
    iterator = worksheet.iter_rows(
        min_row=header_row,
        max_row=header_row,
        values_only=True,
    )
    row = next(iterator, ())
    return trim_trailing_empty(row)


def inspect_modern(
    input_path: Path,
    *,
    header_row: int,
    sample_rows: int,
    sample_columns: int,
    sheet_name: str | None,
    sheet_index: int | None,
    formulas: bool,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        input_path,
        read_only=True,
        data_only=not formulas,
        keep_links=False,
    )
    try:
        all_names = list(workbook.sheetnames)
        if sheet_name is not None or sheet_index is not None:
            selected_name, _ = choose_sheet(all_names, sheet_name, sheet_index)
            names = [selected_name]
        else:
            names = all_names

        sheets: list[dict[str, Any]] = []
        for name in names:
            worksheet = workbook[name]
            headers = read_modern_header(worksheet, header_row)
            width = len(headers)
            samples: list[list[Any]] = []
            if sample_rows and width:
                for row in worksheet.iter_rows(
                    min_row=header_row + 1,
                    max_col=width,
                    values_only=True,
                ):
                    if all(value is None for value in row):
                        continue
                    samples.append([json_value(value) for value in row[:sample_columns]])
                    if len(samples) >= sample_rows:
                        break
            sheets.append(
                {
                    "name": name,
                    "reported_max_row": worksheet.max_row,
                    "reported_max_column": worksheet.max_column,
                    "header_row": header_row,
                    "headers": [
                        json_value(value) for value in headers[:sample_columns]
                    ],
                    "headers_truncated": width > sample_columns,
                    "header_issues": header_issues(headers),
                    "sample_rows": samples,
                }
            )
        return {"sheet_names": all_names, "sheets": sheets}
    finally:
        workbook.close()


def xls_cell_value(book: Any, cell: Any) -> Any:
    import xlrd

    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, book.datemode)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, f"#ERROR-{cell.value}")
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    return cell.value


def read_legacy_row(book: Any, sheet: Any, row_index: int) -> list[Any]:
    return [xls_cell_value(book, sheet.cell(row_index, col)) for col in range(sheet.ncols)]


def inspect_legacy(
    input_path: Path,
    *,
    header_row: int,
    sample_rows: int,
    sample_columns: int,
    sheet_name: str | None,
    sheet_index: int | None,
    formulas: bool,
) -> dict[str, Any]:
    if formulas:
        raise SpreadsheetToolError("formula expressions cannot be extracted from legacy .xls files")

    import xlrd

    workbook = xlrd.open_workbook(input_path, on_demand=True)
    try:
        all_names = workbook.sheet_names()
        if sheet_name is not None or sheet_index is not None:
            selected_name, _ = choose_sheet(all_names, sheet_name, sheet_index)
            names = [selected_name]
        else:
            names = all_names

        sheets: list[dict[str, Any]] = []
        for name in names:
            sheet = workbook.sheet_by_name(name)
            raw_headers = (
                trim_trailing_empty(read_legacy_row(workbook, sheet, header_row - 1))
                if header_row <= sheet.nrows
                else []
            )
            samples: list[list[Any]] = []
            if sample_rows:
                for row_index in range(header_row, sheet.nrows):
                    row = read_legacy_row(workbook, sheet, row_index)
                    if all(value is None for value in row):
                        continue
                    samples.append(
                        [json_value(value) for value in row[:sample_columns]]
                    )
                    if len(samples) >= sample_rows:
                        break
            sheets.append(
                {
                    "name": name,
                    "reported_max_row": sheet.nrows,
                    "reported_max_column": sheet.ncols,
                    "header_row": header_row,
                    "headers": [
                        json_value(value) for value in raw_headers[:sample_columns]
                    ],
                    "headers_truncated": len(raw_headers) > sample_columns,
                    "header_issues": header_issues(raw_headers),
                    "sample_rows": samples,
                }
            )
        return {"sheet_names": all_names, "sheets": sheets}
    finally:
        workbook.release_resources()


def inspect_workbook(input_value: Path, args: argparse.Namespace) -> dict[str, Any]:
    started = time.perf_counter()
    input_path = validate_input(input_value)
    reader = inspect_legacy if input_path.suffix.lower() == ".xls" else inspect_modern
    details = reader(
        input_path,
        header_row=args.header_row,
        sample_rows=args.sample_rows,
        sample_columns=args.sample_columns,
        sheet_name=args.sheet,
        sheet_index=args.sheet_index,
        formulas=args.formulas,
    )
    return {
        "tool": "tools/extract-spreadsheet.py",
        "tool_version": TOOL_VERSION,
        "operation": "inspect",
        "input": {
            "file": display_path(input_path),
            "size_bytes": input_path.stat().st_size,
            "sha256": sha256_file(input_path),
        },
        "cell_mode": "formulas" if args.formulas else "cached-values",
        **details,
        "elapsed_seconds": round(time.perf_counter() - started, 3),
    }


def inspect_command(args: argparse.Namespace) -> int:
    report = inspect_workbook(args.input, args)

    if args.report is None:
        print(json.dumps(report, ensure_ascii=False, indent=2))
        return 0

    report_path = validate_output(input_path, args.report, args.overwrite)
    write_json_atomic(report_path, report)
    print(
        json.dumps(
            {"status": "ok", "report": display_path(report_path)}, ensure_ascii=False
        )
    )
    return 0


def collect_workbooks(inputs: Sequence[Path], recursive: bool) -> list[Path]:
    collected: dict[Path, None] = {}
    for candidate in inputs:
        resolved = candidate.expanduser().resolve()
        if resolved.is_file():
            collected[validate_input(resolved)] = None
            continue
        if not resolved.is_dir():
            raise SpreadsheetToolError(f"input path does not exist: {candidate}")
        iterator = resolved.rglob("*") if recursive else resolved.iterdir()
        for path in iterator:
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXCEL_SUFFIXES:
                collected[path.resolve()] = None
    if not collected:
        raise SpreadsheetToolError("no supported Excel workbooks were found")
    return sorted(collected, key=lambda path: display_path(path).casefold())


def validate_batch_report(
    input_paths: Sequence[Path], report_path: Path, overwrite: bool
) -> Path:
    resolved = report_path.expanduser().resolve()
    if resolved in input_paths:
        raise SpreadsheetToolError("batch report path must not overwrite an input workbook")
    if resolved.exists() and not overwrite:
        raise SpreadsheetToolError(
            f"output already exists: {display_path(resolved)}; pass --overwrite explicitly"
        )
    return resolved


def inspect_many_command(args: argparse.Namespace) -> int:
    started = time.perf_counter()
    input_paths = collect_workbooks(args.inputs, args.recursive)
    reports: list[dict[str, Any]] = []
    failures = 0
    for input_path in input_paths:
        try:
            reports.append({"status": "ok", **inspect_workbook(input_path, args)})
        except Exception as exc:
            failures += 1
            reports.append(
                {
                    "status": "error",
                    "input": {"file": display_path(input_path)},
                    "error": f"{exc.__class__.__name__}: {exc}",
                }
            )
    report = {
        "tool": "tools/extract-spreadsheet.py",
        "tool_version": TOOL_VERSION,
        "operation": "inspect-many",
        "created_at_utc": datetime.now(timezone.utc).isoformat(),
        "workbook_count": len(input_paths),
        "error_count": failures,
        "workbooks": reports,
        "elapsed_seconds": round(time.perf_counter() - started, 3),
    }
    if args.report is None:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        report_path = validate_batch_report(input_paths, args.report, args.overwrite)
        write_json_atomic(report_path, report)
        print(
            json.dumps(
                {
                    "status": "ok" if failures == 0 else "partial",
                    "report": display_path(report_path),
                    "workbooks": len(input_paths),
                    "errors": failures,
                },
                ensure_ascii=False,
            )
        )
    return 1 if failures else 0


def extract_modern(input_path: Path, output_path: Path, args: argparse.Namespace) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(
        input_path,
        read_only=True,
        data_only=not args.formulas,
        keep_links=False,
    )
    try:
        sheet_name, sheet_index = choose_sheet(
            workbook.sheetnames, args.sheet, args.sheet_index
        )
        worksheet = workbook[sheet_name]
        raw_headers = read_modern_header(worksheet, args.header_row)
        selected_indices, headers, warnings = select_columns(
            raw_headers, args.columns, args.repair_headers
        )
        rows = worksheet.iter_rows(
            min_row=args.header_row + 1,
            max_col=len(raw_headers),
            values_only=True,
        )
        statistics = write_rows(
            rows,
            headers,
            selected_indices,
            output_path,
            encoding=args.encoding,
            delimiter="\t" if args.output_suffix == ".tsv" else ",",
            keep_empty_rows=args.keep_empty_rows,
            duplicate_strategy=args.effective_duplicate_check,
            duplicate_memory_rows=args.duplicate_memory_rows,
            max_rows=args.max_rows,
        )
        return {
            "sheet_name": sheet_name,
            "sheet_index": sheet_index + 1,
            "reported_max_row": worksheet.max_row,
            "reported_max_column": worksheet.max_column,
            "header_row": args.header_row,
            "selected_columns": headers,
            "cell_mode": "formulas" if args.formulas else "cached-values",
            "header_warnings": warnings,
            **statistics,
        }
    finally:
        workbook.close()


def extract_legacy(input_path: Path, output_path: Path, args: argparse.Namespace) -> dict[str, Any]:
    if args.formulas:
        raise SpreadsheetToolError("formula expressions cannot be extracted from legacy .xls files")

    import xlrd

    workbook = xlrd.open_workbook(input_path, on_demand=True)
    try:
        sheet_name, sheet_index = choose_sheet(
            workbook.sheet_names(), args.sheet, args.sheet_index
        )
        sheet = workbook.sheet_by_name(sheet_name)
        if args.header_row > sheet.nrows:
            raw_headers: list[Any] = []
        else:
            raw_headers = trim_trailing_empty(
                read_legacy_row(workbook, sheet, args.header_row - 1)
            )
        selected_indices, headers, warnings = select_columns(
            raw_headers, args.columns, args.repair_headers
        )
        rows: Iterator[list[Any]] = (
            read_legacy_row(workbook, sheet, row_index)
            for row_index in range(args.header_row, sheet.nrows)
        )
        statistics = write_rows(
            rows,
            headers,
            selected_indices,
            output_path,
            encoding=args.encoding,
            delimiter="\t" if args.output_suffix == ".tsv" else ",",
            keep_empty_rows=args.keep_empty_rows,
            duplicate_strategy=args.effective_duplicate_check,
            duplicate_memory_rows=args.duplicate_memory_rows,
            max_rows=args.max_rows,
        )
        return {
            "sheet_name": sheet_name,
            "sheet_index": sheet_index + 1,
            "reported_max_row": sheet.nrows,
            "reported_max_column": sheet.ncols,
            "header_row": args.header_row,
            "selected_columns": headers,
            "cell_mode": "cached-values",
            "header_warnings": warnings,
            **statistics,
        }
    finally:
        workbook.release_resources()


def commit_output_pair(
    staged_output: Path,
    output_path: Path,
    staged_audit: Path,
    audit_path: Path,
    overwrite: bool,
) -> None:
    destinations = ((staged_output, output_path), (staged_audit, audit_path))
    backups: dict[Path, Path] = {}
    committed: list[Path] = []
    try:
        for _, destination in destinations:
            if destination.exists():
                if not overwrite:
                    raise SpreadsheetToolError(
                        f"output already exists: {display_path(destination)}; pass --overwrite explicitly"
                    )
                backup = temporary_sibling(destination)
                backup.unlink()
                os.replace(destination, backup)
                backups[destination] = backup
        for staged, destination in destinations:
            os.replace(staged, destination)
            committed.append(destination)
    except Exception:
        for destination in reversed(committed):
            destination.unlink(missing_ok=True)
        for destination, backup in backups.items():
            if backup.exists():
                os.replace(backup, destination)
        raise
    finally:
        for backup in backups.values():
            backup.unlink(missing_ok=True)
        staged_output.unlink(missing_ok=True)
        staged_audit.unlink(missing_ok=True)


def extract_command(args: argparse.Namespace) -> int:
    started = time.perf_counter()
    input_path = validate_input(args.input)
    output_path = validate_output(input_path, args.output, args.overwrite)
    if output_path.suffix.lower() not in {".csv", ".tsv"}:
        raise SpreadsheetToolError("extraction output must use a .csv or .tsv suffix")
    args.output_suffix = output_path.suffix.lower()

    audit_candidate = args.audit_report or output_path.with_name(
        f"{output_path.stem}-audit.json"
    )
    audit_path = validate_output(input_path, audit_candidate, args.overwrite)
    if audit_path == output_path:
        raise SpreadsheetToolError("audit report path must differ from extraction output")

    if args.skip_duplicate_check:
        if args.duplicate_check != "auto":
            raise SpreadsheetToolError(
                "--skip-duplicate-check cannot be combined with a non-default --duplicate-check"
            )
        args.effective_duplicate_check = "off"
    else:
        args.effective_duplicate_check = args.duplicate_check

    reader = extract_legacy if input_path.suffix.lower() == ".xls" else extract_modern
    staged_output = temporary_sibling(output_path)
    staged_audit = temporary_sibling(audit_path)
    try:
        details = reader(input_path, staged_output, args)
        audit = {
            "tool": "tools/extract-spreadsheet.py",
            "tool_version": TOOL_VERSION,
            "operation": "extract",
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "input": {
                "file": display_path(input_path),
                "size_bytes": input_path.stat().st_size,
                "sha256": sha256_file(input_path),
            },
            "output": {
                "file": display_path(output_path),
                "size_bytes": staged_output.stat().st_size,
                "sha256": sha256_file(staged_output),
                "encoding": args.encoding,
                "format": output_path.suffix.lower().lstrip("."),
            },
            **details,
            "elapsed_seconds": round(time.perf_counter() - started, 3),
        }
        staged_audit.write_text(
            json.dumps(audit, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )
        commit_output_pair(
            staged_output,
            output_path,
            staged_audit,
            audit_path,
            args.overwrite,
        )
    finally:
        staged_output.unlink(missing_ok=True)
        staged_audit.unlink(missing_ok=True)
    print(
        json.dumps(
            {
                "status": "ok",
                "output": display_path(output_path),
                "audit_report": display_path(audit_path),
                "rows": details["output_rows"],
                "columns": len(details["selected_columns"]),
            },
            ensure_ascii=False,
        )
    )
    return 0


def add_sheet_selector(parser: argparse.ArgumentParser) -> None:
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--sheet", help="worksheet name; defaults to the first sheet")
    group.add_argument(
        "--sheet-index",
        type=positive_int,
        help="one-based worksheet index; defaults to 1",
    )


def add_inspection_options(parser: argparse.ArgumentParser) -> None:
    add_sheet_selector(parser)
    parser.add_argument(
        "--header-row", type=positive_int, default=1, help="one-based header row"
    )
    parser.add_argument(
        "--sample-rows",
        type=nonnegative_int,
        default=5,
        help="non-empty data rows to sample per selected sheet",
    )
    parser.add_argument(
        "--sample-columns",
        type=positive_int,
        default=20,
        help="maximum columns included in headers and sample rows",
    )
    parser.add_argument(
        "--formulas",
        action="store_true",
        help="show formula expressions instead of cached values where available",
    )
    parser.add_argument("--report", type=Path, help="optional JSON report path")
    parser.add_argument(
        "--overwrite", action="store_true", help="allow replacement of an existing report"
    )


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Inspect Excel workbooks and stream large worksheets to auditable CSV/TSV "
            "without modifying the source file."
        )
    )
    parser.add_argument("--version", action="version", version=TOOL_VERSION)
    commands = parser.add_subparsers(dest="command", required=True)

    inspect_parser = commands.add_parser(
        "inspect", help="list worksheets, headers, dimensions, and bounded sample rows"
    )
    inspect_parser.add_argument("input", type=Path, help="source .xlsx/.xlsm/.xls file")
    add_inspection_options(inspect_parser)
    inspect_parser.set_defaults(func=inspect_command)

    inspect_many_parser = commands.add_parser(
        "inspect-many",
        help="inspect multiple workbooks or directories and produce one combined report",
    )
    inspect_many_parser.add_argument(
        "inputs", nargs="+", type=Path, help="workbook files or directories"
    )
    inspect_many_parser.add_argument(
        "--recursive", action="store_true", help="search input directories recursively"
    )
    add_inspection_options(inspect_many_parser)
    inspect_many_parser.set_defaults(func=inspect_many_command)

    extract_parser = commands.add_parser(
        "extract", help="stream one worksheet to CSV/TSV and create an audit report"
    )
    extract_parser.add_argument("input", type=Path, help="source .xlsx/.xlsm/.xls file")
    extract_parser.add_argument("output", type=Path, help="destination .csv or .tsv file")
    add_sheet_selector(extract_parser)
    extract_parser.add_argument(
        "--header-row", type=positive_int, default=1, help="one-based header row"
    )
    extract_parser.add_argument(
        "--columns",
        nargs="+",
        help="exact normalized header names to extract, in the requested output order",
    )
    extract_parser.add_argument(
        "--repair-headers",
        action="store_true",
        help="name blank headers and suffix duplicates deterministically",
    )
    extract_parser.add_argument(
        "--formulas",
        action="store_true",
        help="export formula expressions instead of cached values where available",
    )
    extract_parser.add_argument(
        "--max-rows",
        type=positive_int,
        help="optional cap on non-empty output rows for trials",
    )
    extract_parser.add_argument(
        "--keep-empty-rows",
        action="store_true",
        help="preserve rows whose selected cells are all empty",
    )
    extract_parser.add_argument(
        "--duplicate-check",
        choices=("auto", "memory", "disk", "off"),
        default="auto",
        help=(
            "duplicate-row strategy; auto switches from memory to an exact SQLite index "
            "after the configured row threshold"
        ),
    )
    extract_parser.add_argument(
        "--duplicate-memory-rows",
        type=positive_int,
        default=DEFAULT_DUPLICATE_MEMORY_ROWS,
        help="unique-row threshold before auto duplicate checking switches to disk",
    )
    extract_parser.add_argument(
        "--skip-duplicate-check",
        action="store_true",
        help="deprecated alias for --duplicate-check off",
    )
    extract_parser.add_argument(
        "--encoding",
        choices=("utf-8", "utf-8-sig"),
        default="utf-8-sig",
        help="text encoding for the output (default: utf-8-sig)",
    )
    extract_parser.add_argument(
        "--audit-report",
        type=Path,
        help="audit JSON path; defaults to <output-stem>-audit.json",
    )
    extract_parser.add_argument(
        "--overwrite",
        action="store_true",
        help="allow replacement of existing output and audit files",
    )
    extract_parser.set_defaults(func=extract_command)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    configure_utf8_stdio()
    parser = build_parser()
    args = parser.parse_args(argv)
    try:
        return int(args.func(args))
    except SpreadsheetToolError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc.__class__.__name__}: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
